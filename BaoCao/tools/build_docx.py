"""Ghép các chương Markdown thành file Word theo template CSE703014.

Cách dùng:  python build_docx.py <output.docx> chuong1.md [chuong2.md ...]
Quy ước Markdown: xem README.md cùng thư mục.
"""
import copy
import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageFont

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(os.path.dirname(ROOT), "..", "CSE703014_Report_Template_VI_V1.docx")
TEMPLATE = os.path.normpath(TEMPLATE)

FONT = "Times New Roman"
BODY_PT = 13
TABLE_PT = 11.5
CODE_PT = 8.5
TEXT_WIDTH_CM = 16.4
MAX_IMG_H_CM = 19.5
LANDSCAPE_WIDTH_CM = 22.8
LANDSCAPE_MAX_H_CM = 14.3
UI_MAX_W_CM = 13.5
UI_MAX_H_CM = 9.0
DIAGRAM_DPI = 150  # ảnh PlantUML render ở 150 dpi: không phóng to quá kích thước gốc

COVER = {
    "team": "",  # đồ án cá nhân: không có Team ID
    "course": "TÊN HỌC PHẦN : ĐỒ ÁN CHUYÊN NGÀNH",
    "title": "TÊN ĐỀ TÀI : XÂY DỰNG HỆ THỐNG QUẢN LÝ THƯ VIỆN CHO NHÀ SÁCH",
    # khối thông tin dạng "nhãn  giá trị"; dòng sinh viên: (họ tên, MSSV, email)
    "info": [
        ("Giảng viên hướng dẫn:", "TS. Nguyễn Thị Vân"),
        ("Sinh viên thực hiện:", None),
        (None, ("Đặng Nhất Nhất", "23010345", "23010345@st.phenikaa-uni.edu.vn")),
        ("Khoá:", "K17 – 2023 – 2027"),
        ("Lớp tín chỉ:", "CSE703096-1-1-26(N03)"),
        ("Chương trình đào tạo:", "Đại học chính quy"),
    ],
    "info_pt": 14,
    "signer": "Đặng Nhất Nhất",
    "date": "Hà Nội, tháng 10 năm 2026",
}
# trang đầu báo cáo: Lời cảm ơn, Lời cam đoan (trước Mục lục) và Lời mở đầu (trước chương 1)
FRONT_BEFORE_TOC = "loi_cam_on.md"
FRONT_BEFORE_CHAPTERS = "loi_mo_dau.md"

# ───────────────────────────── tiện ích XML ─────────────────────────────
def set_run_font(run, size=None, bold=None, italic=None, name=FONT):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def para_shade(p, fill):
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def ensure_style(doc, name, base="Normal"):
    styles = doc.styles
    try:
        return styles[name]
    except KeyError:
        st = styles.add_style(name, 1)
        st.base_style = styles[base]
        return st


# ───────────────────────────── định dạng chung ─────────────────────────────
def setup_styles(doc):
    # mặc định toàn văn bản: TNR 13
    rpr_default = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    sz = rpr_default.find(qn("w:sz"))
    sz.set(qn("w:val"), str(BODY_PT * 2))

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)
    pf.space_before = Pt(0)

    # cỡ chữ tiêu đề theo template (Heading 1, 2 thừa hưởng 14 pt từ mặc định của template);
    # khoảng cách trước/sau giữ nguyên định nghĩa của template
    for name, size in (("Heading 1", 14), ("Heading 2", 14), ("Heading 3", 13)):
        st = doc.styles[name]
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.keep_with_next = True
    # template đặt dòng tiểu mục "cố định 12 pt" làm mất dấu tiếng Việt ở chữ 13 pt → dùng "tối thiểu 12 pt"
    doc.styles["Heading 3"].paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    doc.styles["Heading 1"].paragraph_format.page_break_before = True
    doc.styles["Heading 3"].font.italic = False

    for name in ("CaptionTable", "CaptionFigure"):
        st = ensure_style(doc, name)
        st.font.name = FONT
        st.font.size = Pt(12)
        st.font.bold = False
        st.font.italic = True
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        st.paragraph_format.line_spacing = 1.15
        st.paragraph_format.space_before = Pt(6)
        st.paragraph_format.space_after = Pt(10)
    doc.styles["CaptionTable"].paragraph_format.keep_with_next = True

    # Mục lục: tiêu đề TOC không ngắt trang, các dòng mục lục dễ đọc
    doc.styles["TOC Heading"].paragraph_format.page_break_before = False
    for name, size, bold in (("toc 1", 13, True), ("toc 2", 13, False), ("toc 3", 12.5, False)):
        st = doc.styles[name]
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.italic = False
        st.paragraph_format.line_spacing = 1.3
        st.paragraph_format.space_before = Pt(2 if bold else 0)
        st.paragraph_format.space_after = Pt(0)

    code = ensure_style(doc, "CodeBlock")
    code.font.name = "Consolas"
    code.font.size = Pt(CODE_PT)
    code.paragraph_format.line_spacing = 1.0
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


# ───────────────────────────── inline markdown ─────────────────────────────
INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[CẦN BỔ SUNG[^\]]*\]|\*[^*\s][^*]*\*)")


def add_inline(p, text, size=BODY_PT, bold=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            add_inline_plain(p, part[2:-2], size, True, False)
        elif part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1])
            set_run_font(run, size=size - 1, bold=bold, name="Consolas")
        elif part.startswith("[CẦN BỔ SUNG"):
            run = p.add_run(part)
            set_run_font(run, size=size, bold=bold)
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            add_inline_plain(p, part[1:-1], size, bold, True)
        else:
            add_inline_plain(p, part, size, bold, False)


def add_inline_plain(p, text, size, bold, italic):
    # vẫn tô vàng placeholder nằm trong đoạn đậm/nghiêng
    for seg in re.split(r"(\[CẦN BỔ SUNG[^\]]*\])", text):
        if not seg:
            continue
        run = p.add_run(seg)
        set_run_font(run, size=size, bold=bold, italic=italic)
        if seg.startswith("[CẦN BỔ SUNG"):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


# ───────────────────────────── các khối ─────────────────────────────
class Builder:
    def __init__(self, doc, base_dir):
        self.doc = doc
        self.base_dir = base_dir
        self.captions = {"table": [], "figure": []}
        self.pending_table_caption = None
        self.landscape = False
        self.in_appendix = False

    def set_orientation(self, landscape):
        from docx.enum.section import WD_ORIENT
        if landscape == self.landscape:
            return
        prev = self.doc.sections[-1]
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        w, h = prev.page_width, prev.page_height
        if landscape:
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = max(w, h), min(w, h)
        else:
            sec.orientation = WD_ORIENT.PORTRAIT
            sec.page_width, sec.page_height = min(w, h), max(w, h)
        self.landscape = landscape

    def _no_number(self, p):
        ppr = p._p.get_or_add_pPr()
        numpr = OxmlElement("w:numPr")
        numid = OxmlElement("w:numId")
        numid.set(qn("w:val"), "0")
        numpr.append(numid)
        ppr.append(numpr)

    def heading(self, level, text):
        is_appendix_title = level == 1 and text.strip().lower().startswith("phụ lục")
        is_front = level == 1 and text.isupper()  # LỜI CẢM ƠN, LỜI CAM ĐOAN, LỜI MỞ ĐẦU
        if is_appendix_title:
            self.in_appendix = True
        if not self.in_appendix:
            text = re.sub(r"^\d+(\.\d+)*\.?\s+", "", text.strip())
        else:
            text = text.strip()
        p = self.doc.add_paragraph(style=f"Heading {level}")
        run = p.add_run(text)
        set_run_font(run)
        if self.in_appendix or is_front:
            # Phụ lục và các trang đầu: không đánh số tự động (tiêu đề phụ lục đã có sẵn "A", "A.1"...)
            self._no_number(p)
            if is_appendix_title or is_front:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def signature(self):
        """Khối ký tên lệch phải cuối Lời cảm ơn, Lời cam đoan."""
        for text, bold, italic in ((COVER["date"], False, True), ("Sinh viên thực hiện", True, False),
                                   ("", False, False), ("", False, False), (COVER["signer"], True, False)):
            p = self.doc.add_paragraph(style="Normal")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.left_indent = Cm(8)
            set_run_font(p.add_run(text), bold=bold, italic=italic)

    def subheading(self, text):
        p = self.doc.add_paragraph(style="Normal")
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        add_inline(p, text, bold=True)
        for r in p.runs:
            r.font.italic = True

    def paragraph(self, text):
        if re.match(r"^\[\d+\]\s", text):  # reference_entry: mục tài liệu tham khảo, thụt treo
            p = self.doc.add_paragraph(style="Normal")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.left_indent = Cm(1)
            p.paragraph_format.first_line_indent = Cm(-1)
            p.paragraph_format.line_spacing = 1.3
            add_inline(p, text)
            return
        p = self.doc.add_paragraph(style="Normal")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(p, text)

    def list_item(self, text, level, marker):
        p = self.doc.add_paragraph(style="Normal")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        left = 1.0 + 0.75 * level
        p.paragraph_format.left_indent = Cm(left)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(left))
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(f"{marker}\t")
        set_run_font(run, size=BODY_PT)
        add_inline(p, text)

    def table_caption(self, text):
        p = self.doc.add_paragraph(style="CaptionTable")
        add_inline(p, text, size=12)
        for r in p.runs:
            r.font.italic = True
        self.captions["table"].append(text)

    def table(self, rows):
        header, body = rows[0], rows[1:]
        ncol = len(header)
        t = self.doc.add_table(rows=1 + len(body), cols=ncol)
        t.style = self.doc.styles["Table Grid"]
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblpr = t._tbl.tblPr
        tblw = tblpr.find(qn("w:tblW"))
        if tblw is None:
            tblw = OxmlElement("w:tblW")
            tblpr.append(tblw)
        tblw.set(qn("w:type"), "pct")
        tblw.set(qn("w:w"), "5000")
        # độ rộng cột theo độ dài nội dung, nhưng không hẹp hơn từ dài nhất của cột
        # (tránh mã ngắn như "UC-01" bị ngắt dòng); ~0,22 cm/ký tự ở cỡ chữ bảng + lề ô
        plain = [[re.sub(r"\*\*|`", "", r[i]) if i < len(r) else "" for i in range(ncol)] for r in rows]
        weights = [min(max(max(len(r[i]) for r in plain), 6), 70) for i in range(ncol)]
        prop = [TEXT_WIDTH_CM * w / sum(weights) for w in weights]
        mins = [min(0.4 * TEXT_WIDTH_CM, 0.5 + 0.22 * max(len(w) for r in plain for w in re.split(r"\s+|<br>", r[i])))
                for i in range(ncol)]
        cm = [max(m, p) for m, p in zip(mins, prop)]
        over, slack = sum(cm) - TEXT_WIDTH_CM, sum(c - m for c, m in zip(cm, mins))
        if over > 0 and slack > over:  # lấy bớt phần dư ở các cột còn rộng hơn mức tối thiểu
            cm = [c - over * (c - m) / slack for c, m in zip(cm, mins)]
        elif over > 0:  # không đủ chỗ cho mọi mức tối thiểu: chia đôi giữa mức tối thiểu thu nhỏ và tỉ lệ nội dung
            cm = [(m * TEXT_WIDTH_CM / sum(mins) + p) / 2 for m, p in zip(mins, prop)]
        assert abs(sum(cm) - TEXT_WIDTH_CM) < 0.01, cm
        widths = [Cm(c) for c in cm]
        for ri, row in enumerate([header] + body):
            cells = t.rows[ri].cells
            for ci in range(ncol):
                cell = cells[ci]
                cell.width = widths[ci]
                txt = row[ci] if ci < len(row) else ""
                cell.paragraphs[0].text = ""
                lines = txt.split("<br>")
                for li, line in enumerate(lines):
                    p = cell.paragraphs[0] if li == 0 else cell.add_paragraph()
                    p.paragraph_format.line_spacing = 1.15
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.first_line_indent = Cm(0)
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 else WD_ALIGN_PARAGRAPH.LEFT
                    add_inline(p, line.strip(), size=TABLE_PT, bold=(ri == 0))
                if ri == 0:
                    shade(cell, "D9E2F3")
        # lặp lại dòng tiêu đề khi sang trang
        trpr = t.rows[0]._tr.get_or_add_trPr()
        th = OxmlElement("w:tblHeader")
        th.set(qn("w:val"), "true")
        trpr.append(th)
        # bảng ngắn (như bảng đặc tả use-case) giữ trọn trên một trang, nếu ước lượng vừa một trang
        # shortcut: chiều cao ước lượng theo số ký tự/độ rộng cột (~0,19 cm/ký tự, ~0,5 cm/dòng); bảng nhiều chữ đặc biệt có thể lệch
        est_cm = sum(max(sum(max(1, -(-len(part) // max(1, int((w - 0.4) / 0.19)))) for part in cell.split("<br>"))
                         for cell, w in zip(row, cm)) * 0.5 + 0.15 for row in plain)
        if len(t.rows) <= 12 and est_cm < 19:
            for row in t.rows[:-1]:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.keep_with_next = True
        # khoảng trống sau bảng
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)
        spacer.paragraph_format.line_spacing = 1.0

    def image(self, caption, path):
        full = os.path.join(self.base_dir, path)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.0
        if os.path.exists(full):
            with Image.open(full) as im:
                w, h = im.size
            max_w = LANDSCAPE_WIDTH_CM if self.landscape else TEXT_WIDTH_CM
            max_h = LANDSCAPE_MAX_H_CM if self.landscape else MAX_IMG_H_CM
            width_cm = max_w
            if os.path.basename(path).startswith("h"):  # biểu đồ: tối đa bằng kích thước gốc
                width_cm = min(width_cm, w / DIAGRAM_DPI * 2.54)
            if "ui_report" in path:  # ảnh giao diện: thu nhỏ để ~2 ảnh mỗi trang
                width_cm, max_h = min(width_cm, UI_MAX_W_CM), UI_MAX_H_CM
            if h / w * width_cm > max_h:
                width_cm = max_h * w / h
            p.add_run().add_picture(full, width=Cm(width_cm))
        else:
            run = p.add_run(f"[CẦN BỔ SUNG: chèn ảnh {path}]")
            set_run_font(run)
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        c = self.doc.add_paragraph(style="CaptionFigure")
        add_inline(c, caption, size=12)
        for r in c.runs:
            r.font.italic = True
        self.captions["figure"].append(caption)

    def code(self, lines):
        for i, line in enumerate(lines):
            p = self.doc.add_paragraph(style="CodeBlock")
            para_shade(p, "F2F2F2")
            p.paragraph_format.first_line_indent = Cm(0)
            if i == len(lines) - 1:
                p.paragraph_format.space_after = Pt(8)
            run = p.add_run(line.replace("\t", "    ") or " ")
            set_run_font(run, size=CODE_PT, name="Consolas")

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def parse_table_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in re.split(r"(?<!\\)\|", line)]


def render_markdown(b, text):
    lines = text.splitlines()
    i = 0
    para_buf = []

    def flush():
        if para_buf:
            b.paragraph(" ".join(s.strip() for s in para_buf))
            para_buf.clear()

    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if s.startswith("```"):
            flush()
            j = i + 1
            block = []
            while j < len(lines) and not lines[j].strip().startswith("```"):
                block.append(lines[j])
                j += 1
            b.code(block)
            i = j + 1
            continue
        if s == "":
            flush()
        elif s == "<!-- pagebreak -->":
            flush()
            b.page_break()
        elif s == "<!-- ký tên -->":
            flush()
            b.signature()
        elif s == "<!-- landscape -->":
            flush()
            b.set_orientation(True)
        elif s == "<!-- portrait -->":
            flush()
            b.set_orientation(False)
        elif s.startswith("<!--"):
            flush()
            while i < len(lines) and "-->" not in lines[i]:
                i += 1
        elif m := re.match(r"^####\s+(.*)", s):
            flush()
            b.subheading(m.group(1))
        elif m := re.match(r"^(#{1,3})\s+(.*)", s):
            flush()
            b.heading(len(m.group(1)), m.group(2))
        elif m := re.match(r"^!\[(.+)\]\((.+)\)$", s):
            flush()
            b.image(m.group(1), m.group(2))
        elif re.match(r"^\*\*Bảng \d+[.\d]*:.*\*\*$", s) or re.match(r"^\*\*Bảng [A-Z]\.\d+:.*\*\*$", s):
            flush()
            b.table_caption(s[2:-2])
        elif s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                r = parse_table_row(lines[i])
                if not all(re.match(r"^:?-{2,}:?$", c) for c in r if c):
                    rows.append([c.replace("\\|", "|") for c in r])
                i += 1
            b.table(rows)
            continue
        elif m := re.match(r"^(\s*)[-*]\s+(.*)", line):
            flush()
            level = len(m.group(1)) // 2
            b.list_item(m.group(2), level, "•" if level == 0 else "–")
        elif m := re.match(r"^(\s*)(\d+)\.\s+(.*)", line):
            flush()
            level = len(m.group(1)) // 2
            b.list_item(m.group(3), level, f"{m.group(2)}.")
        else:
            para_buf.append(s)
        i += 1
    flush()


# ───────────────────────────── template ─────────────────────────────
def find_para(doc, startswith, style=None):
    for p in doc.paragraphs:
        if p.text.strip().startswith(startswith) and (style is None or p.style.name == style):
            return p
    raise ValueError(startswith)


def set_para_text(p, text, size=None, bold=None):
    for r in p.runs[1:]:
        r._element.getparent().remove(r._element)
    if not p.runs:
        p.add_run("")
    p.runs[0].text = text
    set_run_font(p.runs[0], size=size, bold=bold)
    if "[CẦN BỔ SUNG" in text:
        p.runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW


def fill_cover(doc):
    p = find_para(doc, "<Team ID>")
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    if COVER["team"]:
        r = p.add_run(COVER["team"])
        set_run_font(r, size=16, bold=True, italic=True)
        if "[CẦN BỔ SUNG" in COVER["team"]:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        r.add_break()
    r2 = p.add_run(COVER["title"])
    set_run_font(r2, size=16, bold=True, italic=True)
    course = find_para(doc, "BÁO CÁO ĐỒ ÁN")
    set_para_text(course, COVER["course"])
    course.runs[0].add_break()  # một dòng trống giữa tên học phần và tên đề tài
    fill_cover_info(doc)
    date_p = find_para(doc, "<date>")
    set_para_text(date_p, COVER["date"], size=10, bold=True)  # template: dòng ngày 10 pt, đậm
    # bớt vài dòng trống để bìa nằm gọn trong 1 trang
    removed = 0
    prev = date_p._p.getprevious()
    while prev is not None and removed < 6:
        cur, prev = prev, prev.getprevious()
        if cur.tag == qn("w:p") and not "".join(cur.itertext()).strip():
            cur.getparent().remove(cur)
            removed += 1
    set_para_text(find_para(doc, "Content", "TOCEntry"), "MỤC LỤC")


def fill_cover_info(doc):
    """Thay hai chỗ điền <DS sinh viên>, <Supervisor> bằng khối "nhãn  giá trị" (nhãn đậm, cột thẳng hàng bằng tab),
    cả khối canh giữa trang."""
    pt = COVER["info_pt"]
    fonts = {b: ImageFont.truetype(f"C:/Windows/Fonts/{'timesbd' if b else 'times'}.ttf", pt * 20) for b in (True, False)}
    cm = lambda s, b=False: fonts[b].getlength(s) / 20 / 72 * 2.54
    student = next(v for k, v in COVER["info"] if k is None)
    val = max(cm(k, True) for k, _ in COVER["info"] if k) + 0.6  # cột giá trị
    mail = val + cm(student[1]) + 0.9                             # cột email
    left = (TEXT_WIDTH_CM - mail - cm(student[2])) / 2
    assert left > 0, "khối thông tin bìa rộng hơn trang"

    first = find_para(doc, "<DS sinh viên")._p
    last = find_para(doc, "<Supervisor>")._p
    for label, value in COVER["info"]:
        p = doc.add_paragraph(style="No Spacing")
        pf = p.paragraph_format
        pf.line_spacing = 1.5
        pf.left_indent = Cm(left + (0.4 if label is None else 0))
        for pos in (val, mail):
            pf.tab_stops.add_tab_stop(Cm(left + pos))
        if label is None:
            set_run_font(p.add_run("\t".join(value)), size=pt)
        else:
            set_run_font(p.add_run(label), size=pt, bold=True)
            if value:
                set_run_font(p.add_run("\t" + value), size=pt)
        first.addprevious(p._p)
    el = first
    while True:  # bỏ hai chỗ điền mẫu và dòng trống giữa chúng
        assert el is first or el is last or not "".join(el.itertext()).strip()
        nxt, done = el.getnext(), el is last
        el.getparent().remove(el)
        el = nxt
        if done:
            break
    for _ in range(2):  # khối mới dài hơn: bớt 2 dòng trống phía dưới để bìa vẫn gọn 1 trang
        if el is not None and el.tag == qn("w:p") and not "".join(el.itertext()).strip():
            nxt = el.getnext()
            el.getparent().remove(el)
            el = nxt


def fill_assignment_table(doc, rows):
    t = doc.tables[0]
    # giữ dòng tiêu đề, xóa các dòng mẫu
    for row in list(t.rows)[1:]:
        row._tr.getparent().remove(row._tr)
    for data in rows:
        cells = t.add_row().cells
        for c, txt in zip(cells, data):
            c.paragraphs[0].text = ""
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
            c.paragraphs[0].paragraph_format.first_line_indent = Cm(0)
            c.paragraphs[0].paragraph_format.line_spacing = 1.15
            add_inline(c.paragraphs[0], txt, size=12)
    for c in t.rows[0].cells:
        for p in c.paragraphs:
            for r in p.runs:
                set_run_font(r, size=12, bold=True)


def strip_template_body(doc):
    """Xóa phần nội dung mẫu từ 'Giới thiệu' trở đi (giữ sectPr cuối)."""
    start = find_para(doc, "Giới thiệu", "Heading 1")._p
    body = start.getparent()
    el = start
    while el is not None:
        nxt = el.getnext()
        if el.tag != qn("w:sectPr"):
            body.remove(el)
        el = nxt


def strip_blank_before(doc, anchor_text="Phân công thực hiện đồ án"):
    """Bỏ các đoạn trống / ngắt trang thủ công của template giữa mục lục và mục Phân công;
    việc sang trang do thuộc tính page_break_before của tiêu đề đảm nhận."""
    anchor = find_para(doc, anchor_text, "Heading 1")._p
    prev = anchor.getprevious()
    while prev is not None and prev.tag == qn("w:p") and not "".join(prev.itertext()).strip():
        cur, prev = prev, prev.getprevious()
        if cur.find(qn("w:pPr")) is not None and cur.find(qn("w:pPr")).find(qn("w:sectPr")) is not None:
            break
        cur.getparent().remove(cur)


CODE_HYPHEN = re.compile(r"\b(UC|FR|NFR|TC|BR|SEC|PERF|VĐ)-(?=\d)")


def no_break_codes(doc):
    """Giữ mã như FR-01, UC-12 và cụm "Bảng 2.1", "Hình 2.1" trên cùng một dòng."""
    for t in list(doc.element.body.iter(qn("w:t"))):
        text = re.sub(r"\b(Bảng|Hình) (?=\d)", "\\1\u00a0", t.text or "")
        parts = CODE_HYPHEN.sub("\\1\x00", text).split("\x00")
        t.text = parts[0]
        if len(parts) == 1:
            continue
        t.set(qn("xml:space"), "preserve")
        run, anchor = t.getparent(), t.getparent()
        for part in parts[1:]:
            # mỗi dấu gạch nối thành <w:noBreakHyphen/> trong một run cùng định dạng
            for child in (OxmlElement("w:noBreakHyphen"), OxmlElement("w:t")):
                new = copy.deepcopy(run)
                for c in list(new):
                    if c.tag != qn("w:rPr"):
                        new.remove(c)
                if child.tag == qn("w:t"):
                    child.text = part
                    child.set(qn("xml:space"), "preserve")
                new.append(child)
                anchor.addnext(new)
                anchor = new


def set_update_fields(doc):
    settings = doc.settings.element
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    settings.append(uf)


def main():
    out = sys.argv[1]
    md_files = sys.argv[2:]
    doc = Document(TEMPLATE)
    setup_styles(doc)
    fill_cover(doc)
    # đồ án cá nhân: một sinh viên thực hiện toàn bộ
    fill_assignment_table(doc, [
        ["23010345", "Đặng Nhất Nhất", "Kỹ thuật phần mềm (KTPM_EL1)",
         "Thực hiện toàn bộ đồ án (đồ án cá nhân): khảo sát hiện trạng và đặc tả yêu cầu; "
         "thiết kế kiến trúc, cơ sở dữ liệu, API và giao diện; cài đặt máy chủ (Node.js, Express, MySQL) "
         "và giao diện (React); tích hợp đăng nhập Google, thanh toán VNPay/MoMo, trợ lý AI; "
         "đóng gói Docker; kiểm thử và viết báo cáo"],
    ])
    strip_template_body(doc)
    strip_blank_before(doc)
    b = Builder(doc, ROOT)
    toc_title = find_para(doc, "MỤC LỤC", "TOCEntry")
    body = doc.element.body
    start = len(body) - 1  # các đoạn mới được chèn ngay trước sectPr cuối
    with open(os.path.join(ROOT, FRONT_BEFORE_TOC), encoding="utf-8") as fh:
        render_markdown(b, fh.read())
    for el in list(body)[start:-1]:  # chuyển Lời cảm ơn, Lời cam đoan lên trước Mục lục
        toc_title._p.addprevious(el)
    toc_title.paragraph_format.page_break_before = True
    for f in [FRONT_BEFORE_CHAPTERS] + md_files:
        with open(os.path.join(ROOT, f), encoding="utf-8") as fh:
            render_markdown(b, fh.read())
    no_break_codes(doc)
    if os.environ.get("UPDATE_FIELDS_ON_OPEN"):
        set_update_fields(doc)
    doc.save(out)
    print(f"OK {out}: {len(b.captions['table'])} bảng, {len(b.captions['figure'])} hình")


if __name__ == "__main__":
    main()
